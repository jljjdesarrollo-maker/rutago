#!/usr/bin/env python3
"""Exporta TODOS los precios intermedios del sistema RutaGo a un XLSX.
Intermedios = claves que contienen '→' en el nombre.
"""

import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'skills/xlsx'))
from templates.base import *

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ─── Datos: extraer intermedios del TS ───
# Todos los intermedios con '→' organizados por ruta/origen

intermedios = {
    'Troncal Loja→Vilcabamba (IDA)': [
        ('Rumi→Vilc', 'Rumizhitana → Vilcabamba', 1.50, 0.75),
        ('Rumi→Mal', 'Rumizhitana → Malacatos', 1.10, 0.55),
        ('Rumi→Prv', 'Rumizhitana → Porvenir', 0.75, 0.40),
        ('Nango→Vilc', 'Nangora → Vilcabamba', 1.25, 0.65),
        ('Nango→Mal', 'Nangora → Malacatos', 0.75, 0.40),
        ('Nango→Carar', 'Nangora → Cararango', 1.00, 0.50),
        ('Land→Vilc', 'Landangui → Vilcabamba', 1.10, 0.55),
        ('Land→Mal', 'Landangui → Malacatos', 0.75, 0.40),
        ('Mal→Vilc', 'Malacatos → Vilcabamba', 1.10, 0.55),
        ('T.Leguas→Mal', 'Tres Leguas → Malacatos', 1.10, 0.55),
        ('Taxiche→Vilc', 'Taxiche → Vilcabamba', 0.75, 0.40),
        ('Cararango→Vilc', 'Cararango → Vilcabamba', 0.75, 0.40),
        ('S.Pedro→Vilc', 'San Pedro → Vilcabamba', 0.75, 0.40),
    ],
    'Troncal Malacatos→Loja (VUELTA)': [
        ('Mal→Rumi', 'Malacatos → Rumizhitana', 1.10, 0.60),
        ('Mal→Nango', 'Malacatos → Nangora', 0.75, 0.40),
        ('Mal→Porv', 'Malacatos → Porvenir', 0.75, 0.40),
    ],
    'El Tambo IDA (desde Malacatos)': [
        ('Mal→Ceibop', 'Malacatos → Ceibopamba', 0.75, 0.40),
        ('Mal→Trinidad', 'Malacatos → Trinidad', 0.75, 0.40),
        ('Mal→S.Jose', 'Malacatos → San José', 0.75, 0.40),
        ('Mal→StoDom', 'Malacatos → Sto. Domingo', 1.00, 0.50),
        ('Mal→N.Dulce', 'Malacatos → Naranjo Dulce', 1.25, 0.65),
        ('Mal→Zhotahu', 'Malacatos → Zhotahuayco', 1.50, 0.75),
        ('Mal→LaMerc', 'Malacatos → La Merced', 1.75, 0.90),
        ('Mal→S.Agust', 'Malacatos → San Agustín', 1.75, 0.90),
        ('Mal→LaEra', 'Malacatos → La Era', 2.00, 1.00),
        ('Mal→LaCap', 'Malacatos → La Capilla', 2.25, 1.15),
        ('Mal→S.Bern', 'Malacatos → San Bernardo', 2.25, 1.15),
        ('Mal→ElTambo', 'Malacatos → El Tambo', 2.25, 1.15),
    ],
    'El Tambo VUELTA (zona El Tambo)': [
        ('Era→Merc', 'La Era → La Merced', 0.75, 0.40),
        ('Era→Malac', 'La Era → Malacatos', 2.00, 1.00),
        ('Merc→Ceibop', 'La Merced → Ceibopamba', 1.50, 0.75),
        ('Merc→Malac', 'La Merced → Malacatos', 1.75, 0.90),
        ('Merc→Land', 'La Merced → Landangui', 2.00, 1.00),
        ('Zhot→Malac', 'Zhotahuayco → Malacatos', 1.50, 0.75),
        ('Zhot→Merc', 'Zhotahuayco → La Merced', 0.75, 0.40),
        ('Zhot→Ceibop', 'Zhotahuayco → Ceibopamba', 1.00, 0.50),
    ],
    'El Tambo VUELTA (Malacatos→Loja)': [
        ('Mal→Peña', 'Malacatos → La Peña', 0.75, 0.40),
        ('Mal→Chorri', 'Malacatos → Chorrillos', 0.75, 0.40),
        ('Mal→Nango2', 'Malacatos → Nangora (vta)', 0.75, 0.40),
        ('Mal→Porv2', 'Malacatos → Porvenir (vta)', 0.75, 0.40),
        ('Mal→T.Leguas', 'Malacatos → Tres Leguas', 1.10, 0.55),
        ('Mal→P.Nuevo', 'Malacatos → Pueblo Nuevo', 1.10, 0.55),
        ('Mal→Rumi2', 'Malacatos → Rumizhitana (vta)', 1.10, 0.55),
        ('Mal→Caja', 'Malacatos → Cajánuma', 1.50, 0.75),
        ('Mal→D.Puen', 'Malacatos → Dos Puentes', 1.50, 0.75),
    ],
    'La Elvira IDA (desde Malacatos)': [
        ('Mal→Cucan', 'Malacatos → Cucanama', 1.60, 0.80),
        ('Mal→Lind', 'Malacatos → Linderos', 2.00, 1.00),
        ('Mal→Santo', 'Malacatos → Santorum', 2.00, 1.00),
        ('Mal→Solan', 'Malacatos → Solanda', 2.00, 1.00),
        ('Mal→Moyoc', 'Malacatos → Moyococha', 2.00, 1.00),
        ('Mal→Tumia', 'Malacatos → Tumianuma', 2.50, 1.25),
        ('Mal→Quina', 'Malacatos → Quinara', 2.50, 1.25),
        ('Mal→Comun', 'Malacatos → Comunidades', 2.50, 1.25),
        ('Mal→Elvira', 'Malacatos → La Elvira', 3.00, 1.50),
    ],
    'La Elvira IDA (desde Vilcabamba)': [
        ('Vilc→Cucan', 'Vilcabamba → Cucanama', 0.75, 0.40),
        ('Vilc→Lind', 'Vilcabamba → Linderos', 1.10, 0.55),
        ('Vilc→Santo', 'Vilcabamba → Santorum', 1.50, 0.65),
        ('Vilc→Solan', 'Vilcabamba → Solanda', 1.50, 0.65),
        ('Vilc→Moyoc', 'Vilcabamba → Moyococha', 1.50, 0.65),
        ('Vilc→Tumia', 'Vilcabamba → Tumianuma', 2.00, 1.00),
        ('Vilc→Quina', 'Vilcabamba → Quinara', 2.00, 1.00),
        ('Vilc→Comun', 'Vilcabamba → Comunidades', 2.00, 1.00),
        ('Vilc→Elvira', 'Vilcabamba → La Elvira', 2.40, 1.20),
    ],
}

# ─── Build XLSX ───
wb = Workbook()
ws = wb.active
ws.title = "Intermedios"

output_path = '/home/z/my-project/download/Precios_Intermedios_RutaGo.xlsx'

# Setup
last_col = 6  # B to F
setup_sheet(ws, title='Precios Intermedios RutaGo — Todos los tramos', last_col=last_col)

# Column widths
ws.column_dimensions['A'].width = 3
ws.column_dimensions['B'].width = 18
ws.column_dimensions['C'].width = 36
ws.column_dimensions['D'].width = 14
ws.column_dimensions['E'].width = 14
ws.column_dimensions['F'].width = 30

# Headers row 4
headers = ['Clave', 'Tramo', 'Normal', 'Media', 'Ruta / Sección']
for col_idx, h in enumerate(headers, start=2):
    ws.cell(row=4, column=col_idx, value=h)
style_header_row(ws, row_num=4, col_start=2, col_end=last_col)

# Data rows
row = 5
for section, items in intermedios.items():
    # Section title row
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=last_col)
    cell = ws.cell(row=row, column=2, value=section)
    cell.font = font_subheader()
    cell.alignment = Alignment(horizontal='left', vertical='center')
    cell.fill = PatternFill('solid', fgColor=SECONDARY)
    for c in range(3, last_col + 1):
        ws.cell(row=row, column=c).fill = PatternFill('solid', fgColor=SECONDARY)
    ws.row_dimensions[row].height = 24
    row += 1

    for i, (clave, tramo, normal, media) in enumerate(items):
        ws.cell(row=row, column=2, value=clave)
        ws.cell(row=row, column=3, value=tramo)
        ws.cell(row=row, column=4, value=normal)
        ws.cell(row=row, column=5, value=media)
        ws.cell(row=row, column=6, value='')
        
        style_data_row(ws, row_num=row, col_start=2, col_end=last_col, row_index=i)
        
        # Format prices
        ws.cell(row=row, column=4).number_format = '$#,##0.00'
        ws.cell(row=row, column=5).number_format = '$#,##0.00'
        ws.cell(row=row, column=4).alignment = align_number()
        ws.cell(row=row, column=5).alignment = align_number()
        ws.cell(row=row, column=2).alignment = align_text()
        ws.cell(row=row, column=3).alignment = align_text()
        
        row += 1

    row += 1  # blank row between sections

# Summary row
total_intermedios = sum(len(v) for v in intermedios.values())
ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
cell = ws.cell(row=row, column=2, value=f'Total intermedios: {total_intermedios}')
cell.font = font_subheader()
cell.alignment = align_text()
cell.fill = PatternFill('solid', fgColor=SECONDARY)
ws.cell(row=row, column=4, value=f'{len(intermedios)} secciones')
ws.cell(row=row, column=4).font = font_subheader()
ws.cell(row=row, column=4).fill = PatternFill('solid', fgColor=SECONDARY)
for c in range(5, last_col + 1):
    ws.cell(row=row, column=c).fill = PatternFill('solid', fgColor=SECONDARY)

wb.properties.creator = 'Z.ai'
wb.save(output_path)
print(f'XLSX exportado: {output_path}')
print(f'Total intermedios: {total_intermedios} en {len(intermedios)} secciones')
