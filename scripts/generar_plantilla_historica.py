"""
Genera plantilla Excel para carga histórica de arqueos VT.
Datos de ejemplo reales extraídos de la libreta del 26-08-2026.
"""
import sys, os
sys.path.insert(0, '/home/z/my-project/skills/xlsx')
from templates.base import *
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Border, Side, Alignment
from openpyxl.utils import get_column_letter

wb = Workbook()

# ============================================================
# HOJA 1: Viajes (frecuencias)
# ============================================================
ws1 = wb.active
ws1.title = "Viajes"
ws1.sheet_view.showGridLines = False

# Colores RutaGo
RG_PRIMARY = "912D26"
RG_LIGHT = "FEF2F1"
RG_GREEN = "1B7D46"
RG_GRAY = "F3F4F6"

# Columnas: A=Fecha, B=VT, C=Ayudante, D=KM, E=Hora, F=Dirección, G=Recaudado, H=Caja Común, I=Total
headers = ["Fecha", "VT", "Ayudante", "KM", "Hora", "Dirección", "Recaudado", "Caja Común", "Total"]
col_widths = [14, 10, 16, 12, 8, 12, 14, 14, 14]

for i, (h, w) in enumerate(zip(headers, col_widths), 1):
    col = get_column_letter(i)
    ws1.column_dimensions[col].width = w
    cell = ws1.cell(row=1, column=i, value=h)
    cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=RG_PRIMARY)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(bottom=Side(style="thin", color=RG_PRIMARY))

ws1.row_dimensions[1].height = 28

# Datos reales del 26-08-2026 (VERIFICAR con la libreta original)
viajes = [
    ["2026-08-26", "VT-001", "B. López", "885/160", "10:15", "IDA",    8.00,  20.00],
    ["2026-08-26", "VT-001", "B. López", "885/160", "14:15", "IDA",   55.00,  13.15],
    ["2026-08-26", "VT-001", "B. López", "885/160", "19:45", "IDA",   21.75,   3.40],
    ["2026-08-26", "VT-001", "B. López", "885/160", "08:20", "VUELTA", 25.00,   0.00],
    ["2026-08-26", "VT-001", "B. López", "885/160", "16:30", "VUELTA", 75.00,   0.00],
    ["2026-08-26", "VT-001", "B. López", "885/160", "19:15", "VUELTA", 35.50,   0.00],
]

for idx, row_data in enumerate(viajes):
    row = idx + 2
    for col, val in enumerate(row_data, 1):
        cell = ws1.cell(row=row, column=col, value=val)
        cell.font = Font(name=FONT_NAME, size=10, color=NEUTRAL_900)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        # Alternating rows
        if idx % 2 == 0:
            cell.fill = PatternFill("solid", fgColor=RG_LIGHT)
        else:
            cell.fill = PatternFill("solid", fgColor="FFFFFF")
        cell.border = Border(bottom=Side(style="hair", color="E0E0E0"))
    # Columna I = fórmula Total = G + H
    total_cell = ws1.cell(row=row, column=9)
    total_cell.value = f"=G{row}+H{row}"
    total_cell.font = Font(name=FONT_NAME, size=10, bold=True, color=RG_PRIMARY)
    total_cell.number_format = "#,##0.00"
    # Formato numérico para G y H
    ws1.cell(row=row, column=7).number_format = "#,##0.00"
    ws1.cell(row=row, column=8).number_format = "#,##0.00"
    ws1.row_dimensions[row].height = 22

# Fila de totales
total_row = len(viajes) + 2
ws1.cell(row=total_row, column=5, value="TOTALES").font = Font(name=FONT_NAME, size=10, bold=True, color=RG_PRIMARY)
ws1.cell(row=total_row, column=5).alignment = Alignment(horizontal="right", vertical="center")
for col in [7, 8, 9]:
    cell = ws1.cell(row=total_row, column=col)
    col_letter = get_column_letter(col)
    cell.value = f"=SUM({col_letter}2:{col_letter}{total_row-1})"
    cell.font = Font(name=FONT_NAME, size=10, bold=True, color=RG_PRIMARY)
    cell.number_format = "#,##0.00"
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = Border(top=Side(style="medium", color=RG_PRIMARY))
ws1.row_dimensions[total_row].height = 26

# ============================================================
# HOJA 2: Gastos
# ============================================================
ws2 = wb.create_sheet("Gastos")
ws2.sheet_view.showGridLines = False

gastos_headers = ["Fecha", "Concepto", "Monto"]
gastos_widths = [14, 20, 14]

for i, (h, w) in enumerate(zip(gastos_headers, gastos_widths), 1):
    col = get_column_letter(i)
    ws2.column_dimensions[col].width = w
    cell = ws2.cell(row=1, column=i, value=h)
    cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=RG_PRIMARY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = Border(bottom=Side(style="thin", color=RG_PRIMARY))
ws2.row_dimensions[1].height = 28

gastos = [
    ["2026-08-26", "Chofer", 30.00],
    ["2026-08-26", "Ayudante", 20.00],
    ["2026-08-26", "Diesel", 100.75],
    ["2026-08-26", "Plan Renova", 68.00],
    ["2026-08-26", "Fundas", 6.50],
]

for idx, row_data in enumerate(gastos):
    row = idx + 2
    for col, val in enumerate(row_data, 1):
        cell = ws2.cell(row=row, column=col, value=val)
        cell.font = Font(name=FONT_NAME, size=10, color=NEUTRAL_900)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        if idx % 2 == 0:
            cell.fill = PatternFill("solid", fgColor=RG_LIGHT)
        else:
            cell.fill = PatternFill("solid", fgColor="FFFFFF")
        cell.border = Border(bottom=Side(style="hair", color="E0E0E0"))
    ws2.cell(row=row, column=3).number_format = "#,##0.00"
    ws2.row_dimensions[row].height = 22

# Total gastos
total_row_g = len(gastos) + 2
ws2.cell(row=total_row_g, column=2, value="TOTAL GASTOS").font = Font(name=FONT_NAME, size=10, bold=True, color=RG_PRIMARY)
ws2.cell(row=total_row_g, column=2).alignment = Alignment(horizontal="right", vertical="center")
ws2.cell(row=total_row_g, column=3, value=f"=SUM(C2:C{total_row_g-1})")
ws2.cell(row=total_row_g, column=3).font = Font(name=FONT_NAME, size=10, bold=True, color=RG_PRIMARY)
ws2.cell(row=total_row_g, column=3).number_format = "#,##0.00"
ws2.cell(row=total_row_g, column=3).alignment = Alignment(horizontal="center", vertical="center")
ws2.cell(row=total_row_g, column=3).border = Border(top=Side(style="medium", color=RG_PRIMARY))
ws2.row_dimensions[total_row_g].height = 26

# ============================================================
# HOJA 3: Cierre
# ============================================================
ws3 = wb.create_sheet("Cierre")
ws3.sheet_view.showGridLines = False

cierre_headers = ["Fecha", "VT", "Tickets", "Sobrante"]
cierre_widths = [14, 10, 14, 14]

for i, (h, w) in enumerate(zip(cierre_headers, cierre_widths), 1):
    col = get_column_letter(i)
    ws3.column_dimensions[col].width = w
    cell = ws3.cell(row=1, column=i, value=h)
    cell.font = Font(name=FONT_NAME, size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor=RG_PRIMARY)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = Border(bottom=Side(style="thin", color=RG_PRIMARY))
ws3.row_dimensions[1].height = 28

cierre = [
    ["2026-08-26", "VT-001", 4.60, 0.00],
]

for idx, row_data in enumerate(cierre):
    row = idx + 2
    for col, val in enumerate(row_data, 1):
        cell = ws3.cell(row=row, column=col, value=val)
        cell.font = Font(name=FONT_NAME, size=10, color=NEUTRAL_900)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(bottom=Side(style="hair", color="E0E0E0"))
    ws3.cell(row=row, column=3).number_format = "#,##0.00"
    ws3.cell(row=row, column=4).number_format = "#,##0.00"
    ws3.row_dimensions[row].height = 22

# ============================================================
# HOJA 4: Instrucciones
# ============================================================
ws4 = wb.create_sheet("Instrucciones")
ws4.sheet_view.showGridLines = False
ws4.column_dimensions["A"].width = 3
ws4.column_dimensions["B"].width = 80

instrucciones = [
    ("PLANTILLA DE CARGA HISTORICA - RUTAGO", True, RG_PRIMARY, 14),
    ("", False, None, 8),
    ("REGLAS:", True, RG_PRIMARY, 11),
    ("1. La Fecha vincula las 3 hojas (Viajes + Gastos + Cierre). Mismo dia = mismo arqueo.", False, None, 10),
    ("2. Columna 'Total' en Viajes se calcula automaticamente (Recaudado + Caja Comun). No editar.", False, None, 10),
    ("3. Desde JUNIO 2026 en IDA: Caja Común = lo vendido en oficina Loja. Total = Recaudado + Caja Común.", False, None, 10),
    ("4. Desde JUNIO 2026 en VUELTA: Caja Común = 0.00. Solo importa el Recaudado del ayudante.", False, None, 10),
    ("5. ANTES de junio 2026 (IDA y VUELTA): Caja Común = 0.00. Solo importa el Recaudado.", False, None, 10),
    ("6. Los datos se guardan en la misma BD que el arqueo general de la app.", False, None, 10),
    ("", False, None, 8),
    ("HOJAS:", True, RG_PRIMARY, 11),
    ("- Viajes: Una fila por frecuencia. Fecha, VT, Ayudante, KM, Hora, Dirección, Recaudado, Caja Común.", False, None, 10),
    ("- Gastos: Una fila por gasto del día. Fecha, Concepto, Monto.", False, None, 10),
    ("- Cierre: Una fila por día. Fecha, VT, Tickets, Sobrante.", False, None, 10),
    ("", False, None, 8),
    ("IMPORTANTE: Verifica los valores de ejemplo con tu libreta original antes de usarlos como referencia.", True, "C0392B", 10),
]

for idx, (text, bold, color, size) in enumerate(instrucciones):
    row = idx + 1
    cell = ws4.cell(row=row, column=2, value=text)
    if bold:
        cell.font = Font(name=FONT_NAME, size=size, bold=True, color=color or NEUTRAL_900)
    else:
        cell.font = Font(name=FONT_NAME, size=size, color=NEUTRAL_900)
    ws4.row_dimensions[row].height = 20

# Guardar
output_path = "/home/z/my-project/download/plantilla-carga-historica-rutago.xlsx"
wb.save(output_path)
print(f"Archivo generado: {output_path}")
